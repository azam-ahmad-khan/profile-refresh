# /// script
# requires-python = ">=3.9"
# dependencies = ["python-docx>=1.1"]
# ///
"""Build an ATS-friendly resume .docx from a simple Markdown file (format in references/resume.md).

Usage: uv run resume_docx.py resume.md -o Name-Resume-2026-10-18.docx [--style plain|polished]
       [--font Arial] [--size 10] [--accent 1F4E79] [--band EDEDED] [--paper letter|a4] [--margin 0.6]
Single column, real text only, no tables or text boxes. Emails, phone numbers (in the header) and URLs become
clickable links while keeping their full text visible, since some parsers drop hyperlinks. Role headers never split from their first bullet; bullet
lists flow freely across pages, so pages fill instead of leaving gaps.
"""
import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DATE_RE = re.compile(r"((19|20)\d{2}|present|current)", re.I)
EMAIL_RE = r"[\w.+-]+@[\w-]+\.[\w.-]*\w"
# A URL never ends in sentence punctuation, so "see example.org/x." links to example.org/x.
URL_RE = (r"https?://[^\s|,;)]*[^\s|,;).!?:]|(?:www\.)?[a-z0-9-]+(?:\.[a-z0-9-]+)*\."
          r"(?:com|org|net|io|dev|ai|me|co|app|page|site|uk|ae|in|edu|gov|info|bio)(?:/[^\s|,;)]*[^\s|,;).!?:])?")
PHONE_RE = r"\+?\d[\d\s().-]{7,17}\d"


def link_target(text, in_header):
    if re.fullmatch(EMAIL_RE, text):
        return f"mailto:{text}"
    if in_header and re.fullmatch(PHONE_RE, text) and len(re.sub(r"\D", "", text)) >= 9:
        return "tel:" + ("+" if text.strip().startswith("+") else "") + re.sub(r"\D", "", text)
    # Outside the contact header a bare "Socket.io" or "ASP.NET" is a skill, not a link: require a scheme, www or path.
    bare_domain = not re.match(r"(https?://|www\.)", text, re.I) and "/" not in text
    if re.fullmatch(URL_RE, text, re.I) and (in_header or not bare_domain):
        return text if text.lower().startswith("http") else f"https://{text}"
    return None


def add_hyperlink(paragraph, url, text, font, size, bold, color):
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    hyperlink.append(run._element)
    paragraph._p.append(hyperlink)


def add_runs(paragraph, text, font, size, bold=False, color=None, link_phones=False):
    # **bold** spans become bold runs; emails, URLs and (in the header) phone numbers become hyperlinks.
    pattern = f"({EMAIL_RE}|{URL_RE}" + (f"|{PHONE_RE})" if link_phones else ")")
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        for piece in re.split(pattern, part, flags=re.I):
            if not piece:
                continue
            target = link_target(piece, link_phones)
            if target:
                add_hyperlink(paragraph, target, piece, font, size, bold or (i % 2 == 1), color)
                continue
            run = paragraph.add_run(piece)
            run.bold = bold or (i % 2 == 1)
            run.font.name = font
            run.font.size = Pt(size)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def shade(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def bottom_border(paragraph, color="000000"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(key), value)
    borders.append(bottom)
    ppr.append(borders)


def spacing(paragraph, before=0, after=0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def build(md_path, out_path, args):
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    polished = args.style == "polished"
    accent = args.accent if polished else None
    doc = Document()
    section = doc.sections[0]
    if args.paper == "a4":
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(args.margin))
    text_width = section.page_width - section.left_margin - section.right_margin

    in_header = False
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            p = add_runs(doc.add_paragraph(), line[2:].strip(), args.font, args.size + 9, bold=True, color=accent)
            spacing(p, 0, 2)
            in_header = True
        elif line.startswith("## "):
            in_header = False
            p = add_runs(doc.add_paragraph(), line[3:].strip().upper(), args.font, args.size + 1.5, bold=True,
                         color=accent)
            spacing(p, 8, 3)
            p.paragraph_format.keep_with_next = True
            if polished:
                shade(p, args.band)
            else:
                bottom_border(p)
        elif line.startswith("### "):
            parts = [x.strip() for x in line[4:].split("|")]
            dates = parts.pop() if len(parts) > 1 and DATE_RE.search(parts[-1]) else ""
            title, rest = parts[0], ", ".join(parts[1:])
            p = doc.add_paragraph()
            add_runs(p, title, args.font, args.size + 0.5, bold=True, color=accent)
            if rest:
                add_runs(p, f" | {rest}", args.font, args.size + 0.5)
            if dates:
                p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
                add_runs(p, f"\t{dates}", args.font, args.size)
            spacing(p, 5, 1)
            # The role header stays with its first bullet; the bullets themselves may flow onto the next page.
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
        elif re.match(r"^\s*[-*•] ", line):
            p = add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•] ", "", line), args.font, args.size)
            spacing(p, 0, 1)
        else:
            size = args.size + 1 if in_header and len(doc.paragraphs) == 1 else args.size
            p = add_runs(doc.add_paragraph(), line.strip(), args.font, size, bold=in_header and len(doc.paragraphs) == 1,
                         link_phones=in_header)
            spacing(p, 0, 2)
    doc.core_properties.title = Path(out_path).stem
    doc.save(out_path)
    print(f"wrote {out_path}")


def to_pdf(docx_path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("PDF: no LibreOffice found. Open the .docx in Word, Pages or Google Docs and export it as PDF "
              "(File → Save As / Export → PDF), keeping the same file name.")
        return
    out_dir = str(Path(docx_path).resolve().parent)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path], check=True,
                   stdout=subprocess.DEVNULL)
    print(f"wrote {Path(docx_path).with_suffix('.pdf')}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--style", choices=["plain", "polished"], default="plain")
    ap.add_argument("--font", default="Arial")
    ap.add_argument("--size", type=float, default=10)
    ap.add_argument("--accent", default="1F4E79", help="hex colour for name, headings, titles (polished only)")
    ap.add_argument("--band", default="EDEDED", help="hex fill behind section headings (polished only)")
    ap.add_argument("--paper", choices=["letter", "a4"], default="letter")
    ap.add_argument("--margin", type=float, default=0.6, help="inches")
    ap.add_argument("--pdf", action="store_true", help="also convert to PDF if LibreOffice is installed")
    args = ap.parse_args()
    for hex_value in (args.accent, args.band):
        if not re.fullmatch(r"[0-9A-Fa-f]{6}", hex_value):
            sys.exit(f"colour must be 6 hex digits, got {hex_value!r}")
    build(args.markdown, args.out, args)
    if args.pdf:
        to_pdf(args.out)


if __name__ == "__main__":
    main()
